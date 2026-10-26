"""Local project records and a generic storyboard table; no generation engine."""
from __future__ import annotations

from contextlib import contextmanager
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import uuid

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

VERSION = "0.1.0"
MAX_SHOTS = 1000
MAX_SOURCE_CHARS = 200_000
FIELDS = ("shot_id", "scene", "visual", "dialogue", "duration_seconds", "source_id", "notes")
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def text(value, label, maximum, *, required=False):
    if not isinstance(value, str) or len(value) > maximum or INVALID_XML.search(value):
        raise ValueError(f"{label} must be valid text of at most {maximum} characters")
    if required and not value.strip():
        raise ValueError(f"{label} cannot be blank")
    return value  # Preserve whitespace, punctuation, and line endings.


def identifier(value):
    if not isinstance(value, str) or not re.fullmatch(r"[0-9a-f]{32}", value):
        raise ValueError("Invalid record ID")
    return value


def validate_shots(shots):
    if not isinstance(shots, list) or not 1 <= len(shots) <= MAX_SHOTS:
        raise ValueError(f"Supply 1 to {MAX_SHOTS} shots; an empty replacement is refused")
    result, seen = [], set()
    for item in shots:
        if not isinstance(item, dict) or set(item) - set(FIELDS):
            raise ValueError("Shot contains unknown fields; use the basic schema")
        shot_id = item.get("shot_id")
        if not isinstance(shot_id, str) or not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{0,63}", shot_id):
            raise ValueError("shot_id must be 1-64 ASCII letters, digits, dots, underscores or hyphens")
        if shot_id in seen:
            raise ValueError("Duplicate shot_id")
        seen.add(shot_id)
        duration = item.get("duration_seconds")
        if duration is not None and (type(duration) not in (int, float) or not 0 < duration <= 86400 or not math.isfinite(duration)):
            raise ValueError("duration_seconds must be null or a finite positive number up to 86400")
        source_id = item.get("source_id")
        if source_id is not None:
            identifier(source_id)
        result.append({"shot_id": shot_id,
                       "scene": text(item.get("scene", ""), "scene", 120),
                       "visual": text(item.get("visual", ""), "visual", 12000, required=True),
                       "dialogue": text(item.get("dialogue", ""), "dialogue", 12000),
                       "duration_seconds": duration, "source_id": source_id,
                       "notes": text(item.get("notes", ""), "notes", 12000)})
    if len(json.dumps(result, ensure_ascii=False).encode("utf-8")) > 2 * 1024 * 1024:
        raise ValueError("Storyboard payload must not exceed 2 MiB")
    return result


class Store:
    def __init__(self, root: Path):
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.database = self.root / "shortdrama-basic.sqlite3"
        if self.database.is_symlink():
            raise ValueError("Database must not be a symbolic link")
        with self.connection(write=True) as db:
            schema_version = db.execute("PRAGMA user_version").fetchone()[0]
            if schema_version not in (0, 1):
                raise ValueError("Unsupported database schema; keep the existing data intact")
            tables = db.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
            if schema_version == 0 and tables:
                raise ValueError("Refusing to initialize an unrelated database")
            db.execute("""CREATE TABLE IF NOT EXISTS projects (
                id TEXT PRIMARY KEY, name TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%fZ','now')))""")
            db.execute("""CREATE TABLE IF NOT EXISTS sources (
                id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id),
                title TEXT NOT NULL, body TEXT NOT NULL)""")
            db.execute("""CREATE TABLE IF NOT EXISTS shots (
                project_id TEXT NOT NULL REFERENCES projects(id), position INTEGER NOT NULL,
                body TEXT NOT NULL, PRIMARY KEY(project_id, position))""")
            db.execute("PRAGMA user_version=1")

    @contextmanager
    def connection(self, *, write=False):
        db = sqlite3.connect(self.database, timeout=30)
        db.row_factory = sqlite3.Row
        try:
            db.execute("PRAGMA foreign_keys=ON")
            db.execute("BEGIN IMMEDIATE" if write else "BEGIN")
            yield db
            db.commit()
        except BaseException:
            db.rollback()
            raise
        finally:
            db.close()

    @staticmethod
    def project(db, project_id):
        row = db.execute("SELECT * FROM projects WHERE id=?", (identifier(project_id),)).fetchone()
        if row is None:
            raise ValueError("Project not found")
        return dict(row)

    def create_project(self, name):
        name = text(name, "name", 120, required=True)
        project_id = uuid.uuid4().hex
        with self.connection(write=True) as db:
            db.execute("INSERT INTO projects(id,name) VALUES (?,?)", (project_id, name))
            return self.project(db, project_id)

    def list_projects(self, limit=50, offset=0):
        if type(limit) is not int or not 1 <= limit <= 100 or type(offset) is not int or offset < 0:
            raise ValueError("limit must be 1-100 and offset must be nonnegative integers")
        with self.connection() as db:
            items = [dict(row) for row in db.execute("SELECT * FROM projects ORDER BY created_at,id LIMIT ? OFFSET ?", (limit, offset))]
            return {"projects": items, "total": db.execute("SELECT count(*) FROM projects").fetchone()[0]}

    def import_source_text(self, project_id, title, body):
        title = text(title, "title", 120, required=True)
        body = text(body, "text", MAX_SOURCE_CHARS, required=True)
        source_id = uuid.uuid4().hex
        with self.connection(write=True) as db:
            self.project(db, project_id)
            db.execute("INSERT INTO sources VALUES (?,?,?,?)", (source_id, project_id, title, body))
            db.execute("UPDATE projects SET revision=revision+1 WHERE id=?", (project_id,))
            return {"source_id": source_id, "title": title, "characters": len(body),
                    "revision": self.project(db, project_id)["revision"]}

    def read_source_text(self, project_id, source_id):
        with self.connection() as db:
            self.project(db, project_id)
            row = db.execute("SELECT id,title,body FROM sources WHERE project_id=? AND id=?", (project_id, identifier(source_id))).fetchone()
            if row is None:
                raise ValueError("Source not found in this project")
            return {"source_id": row["id"], "title": row["title"], "text": row["body"]}

    def get_storyboard(self, project_id):
        with self.connection() as db:
            result = self.project(db, project_id)
            result["sources"] = [dict(row) for row in db.execute("SELECT id,title,length(body) AS characters FROM sources WHERE project_id=? ORDER BY rowid", (project_id,))]
            result["shots"] = [json.loads(row[0]) for row in db.execute("SELECT body FROM shots WHERE project_id=? ORDER BY position", (project_id,))]
            return result

    def save_storyboard(self, project_id, shots, expected_revision):
        shots = validate_shots(shots)
        if type(expected_revision) is not int or expected_revision < 0:
            raise ValueError("expected_revision must be a nonnegative integer")
        with self.connection(write=True) as db:
            project = self.project(db, project_id)
            if project["revision"] != expected_revision:
                raise ValueError("Revision conflict: read the current storyboard and reapply your edits")
            known_sources = {row[0] for row in db.execute("SELECT id FROM sources WHERE project_id=?", (project_id,))}
            if any(shot["source_id"] is not None and shot["source_id"] not in known_sources for shot in shots):
                raise ValueError("Every source_id must belong to this project")
            db.execute("DELETE FROM shots WHERE project_id=?", (project_id,))
            db.executemany("INSERT INTO shots VALUES (?,?,?)", [(project_id, i, json.dumps(shot, ensure_ascii=False)) for i, shot in enumerate(shots)])
            db.execute("UPDATE projects SET revision=revision+1 WHERE id=?", (project_id,))
            return {"project_id": project_id, "revision": expected_revision + 1, "shot_count": len(shots)}

    def export_storyboard(self, project_id):
        project = self.get_storyboard(project_id)  # One consistent database snapshot.
        if not project["shots"]:
            raise ValueError("Save at least one shot before exporting")
        output = self.root / "exports"
        output.mkdir(exist_ok=True)
        if output.is_symlink() or output.resolve().parent != self.root:
            raise ValueError("Export directory must stay inside the configured data directory")
        target = output / f"{project_id}-r{project['revision']}-{uuid.uuid4().hex[:8]}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Storyboard"
        sheet.append(["Order", "Shot ID", "Scene", "Visual", "Dialogue", "Duration (s)", "Source ID", "Notes"])
        for order, shot in enumerate(project["shots"], 1):
            sheet.append([order] + [shot[field] for field in FIELDS])
        metadata = workbook.create_sheet("Project")
        for row in (("Project", project["name"]), ("Project ID", project_id), ("Revision", project["revision"]), ("Edition", f"Shortdrama Basic {VERSION}"), ("Source ID", "Title", "Characters")):
            metadata.append(row)
        for source in project["sources"]:
            metadata.append([source["id"], source["title"], source["characters"]])
        for page in workbook:
            for row in page:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # Treat user text as text, including =formula-looking values.
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            page.freeze_panes = "A2"
            for cell in page[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245C4A")
        for column, width in zip("ABCDEFGH", (8, 18, 20, 48, 42, 16, 36, 32)):
            sheet.column_dimensions[column].width = width
        metadata.column_dimensions["A"].width = 36
        metadata.column_dimensions["B"].width = 52
        metadata.column_dimensions["C"].width = 16
        sheet.auto_filter.ref = sheet.dimensions
        descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=output)
        os.close(descriptor)
        try:
            workbook.save(temporary)
            os.replace(temporary, target)
        finally:
            workbook.close()
            Path(temporary).unlink(missing_ok=True)
        return {"project_id": project_id, "revision": project["revision"],
                "shot_count": len(project["shots"]), "path": str(target), "bytes": target.stat().st_size}
