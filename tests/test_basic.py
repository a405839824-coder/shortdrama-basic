"""Behavior checks for storage, isolation, edits, export and source packaging."""
from concurrent.futures import ThreadPoolExecutor
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
import uuid
import zipfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "server"), str(ROOT / "scripts")]
from basic import Store, MAX_SOURCE_CHARS
from configure_local_mcp import configure
from package_source import PUBLIC_FILES, build


class BasicTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="storyboard basic test ")
        self.addCleanup(self.temporary.cleanup)
        self.store = Store(Path(self.temporary.name) / "data")
        self.project = self.store.create_project("演示项目")
        self.project_id = self.project["id"]
        self.row = {"shot_id": "S01", "visual": "雨停了。", "dialogue": "谢谢你。", "duration_seconds": 5}

    def test_persistence_text_order_and_excel_as_literal_text(self):
        original = "  车站。\r\n她说：‘谢谢你。’\n"
        source = self.store.import_source_text(self.project_id, "=原文示例", original)
        row = {**self.row, "source_id": source["source_id"], "visual": "=1+1", "notes": "+SUM(A1:A2)"}
        self.store.save_storyboard(self.project_id, [row, {**self.row, "shot_id": "S02"}], source["revision"])
        reopened = Store(self.store.root)
        self.assertEqual(original, reopened.read_source_text(self.project_id, source["source_id"])["text"])
        saved = reopened.get_storyboard(self.project_id)
        self.assertEqual(["S01", "S02"], [item["shot_id"] for item in saved["shots"]])
        exported = reopened.export_storyboard(self.project_id)
        workbook = load_workbook(exported["path"], data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(["Storyboard", "Project"], workbook.sheetnames)
        self.assertEqual("=1+1", workbook["Storyboard"]["D2"].value)
        self.assertEqual("s", workbook["Storyboard"]["D2"].data_type)
        self.assertEqual("谢谢你。", workbook["Storyboard"]["E2"].value)
        self.assertEqual("s", workbook["Project"]["B6"].data_type)
        self.assertFalse(any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row))
        self.assertEqual(3, workbook["Storyboard"].max_row)

    def test_stale_edit_and_source_import_cannot_overwrite(self):
        self.store.save_storyboard(self.project_id, [self.row], 0)
        before = self.store.get_storyboard(self.project_id)
        self.store.import_source_text(self.project_id, "补充", "补充文本")
        with self.assertRaisesRegex(ValueError, "Revision conflict"):
            self.store.save_storyboard(self.project_id, [{**self.row, "visual": "过时修改"}], before["revision"])
        self.assertEqual(before["shots"], self.store.get_storyboard(self.project_id)["shots"])

    def test_project_source_isolation(self):
        other = self.store.create_project("另一个项目")["id"]
        source = self.store.import_source_text(other, "独立原文", "另一段文本")
        with self.assertRaises(ValueError):
            self.store.read_source_text(self.project_id, source["source_id"])
        with self.assertRaises(ValueError):
            self.store.save_storyboard(self.project_id, [{**self.row, "source_id": source["source_id"]}], 0)
        self.assertEqual(0, self.store.get_storyboard(self.project_id)["revision"])

    def test_invalid_edits_leave_existing_data_unchanged(self):
        self.store.save_storyboard(self.project_id, [self.row], 0)
        before = self.store.get_storyboard(self.project_id)
        bad_tables = [[], [self.row, self.row], [{**self.row, "unexpected": 1}],
                      [{**self.row, "visual": ""}], [{**self.row, "notes": "bad\x01"}],
                      [{**self.row, "visual": "x" * 12001}], [{**self.row, "shot_id": "../x"}]]
        bad_tables += [[{**self.row, "duration_seconds": value}] for value in (0, -1, float("nan"), float("inf"), True, "5", 10 ** 1000)]
        for table in bad_tables:
            with self.subTest(table=repr(table)[:100]), self.assertRaises(ValueError):
                self.store.save_storyboard(self.project_id, table, before["revision"])
            self.assertEqual(before, self.store.get_storyboard(self.project_id))
        with self.assertRaises(ValueError):
            self.store.save_storyboard(self.project_id, [self.row], True)

    def test_invalid_text_and_ids_do_not_create_records(self):
        for value in ("", "x" * (MAX_SOURCE_CHARS + 1), "bad\x00text", "\ud800"):
            with self.subTest(value_length=len(value)), self.assertRaises(ValueError):
                self.store.import_source_text(self.project_id, "source", value)
        self.assertEqual([], self.store.get_storyboard(self.project_id)["sources"])
        with self.assertRaises(ValueError):
            self.store.get_storyboard("../data")
        with self.assertRaises(ValueError):
            self.store.list_projects(-1)

    def test_concurrent_edits_have_exactly_one_winner(self):
        def save(label):
            try:
                return self.store.save_storyboard(self.project_id, [{**self.row, "visual": label}], 0)
            except ValueError as error:
                return str(error)
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(save, ["修改甲", "修改乙"]))
        self.assertEqual(1, sum(isinstance(result, dict) for result in results))
        self.assertEqual(1, sum(isinstance(result, str) and "Revision conflict" in result for result in results))
        self.assertEqual(1, self.store.get_storyboard(self.project_id)["revision"])

    def test_empty_export_and_nonexistent_project_fail(self):
        with self.assertRaises(ValueError):
            self.store.export_storyboard(self.project_id)
        with self.assertRaises(ValueError):
            self.store.get_storyboard("0" * 32)
        self.assertFalse((self.store.root / "exports").exists())


class PackagingTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="basic package test ")
        self.addCleanup(self.temporary.cleanup)
        self.root = (Path(self.temporary.name) / "source").resolve()
        for name in PUBLIC_FILES:
            target = self.root / name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / name, target)
        configure(self.root, portable=True)
        self.output = Path(self.temporary.name) / "source.zip"

    def test_allowlist_excludes_extra_files_and_rejects_local_config(self):
        canary = uuid.uuid4().hex
        for name in ("data/private.txt", ".git/history", "extra/internal.txt", ".env"):
            target = self.root / name
            target.parent.mkdir(exist_ok=True, parents=True)
            target.write_text(canary, encoding="utf-8")
        report = build(self.root, self.output)
        self.assertEqual(len(PUBLIC_FILES), report["files"])
        with zipfile.ZipFile(self.output) as archive:
            self.assertEqual({f"shortdrama-basic/{name}" for name in PUBLIC_FILES}, set(archive.namelist()))
            self.assertFalse(any(canary.encode() in archive.read(name) for name in archive.namelist()))
        config = json.loads((self.root / ".mcp.json").read_text(encoding="utf-8"))
        config["mcpServers"]["shortdrama-basic"]["cwd"] = str(self.root)
        (self.root / ".mcp.json").write_text(json.dumps(config), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "portable"):
            build(self.root, self.output)

    def test_setup_configuration_and_secret_guard(self):
        with self.assertRaises(ValueError):
            configure(self.root)
        executable = self.root / ".venv" / ("Scripts/python.exe" if sys.platform == "win32" else "bin/python")
        executable.parent.mkdir(parents=True)
        executable.touch()
        configure(self.root)
        config = json.loads((self.root / ".mcp.json").read_text(encoding="utf-8"))["mcpServers"]["shortdrama-basic"]
        self.assertEqual(executable.as_posix(), config["command"])
        self.assertEqual((self.root / "server/main.py").as_posix(), config["args"][0])
        configure(self.root, portable=True)
        with (self.root / "README.md").open("a", encoding="utf-8") as handle:
            handle.write("sk-proj-" + "x" * 30)
        with self.assertRaisesRegex(ValueError, "Credential"):
            build(self.root, self.output)


if __name__ == "__main__":
    unittest.main()
