"""Exercise the installed MCP server using isolated, synthetic local data."""
import asyncio
import json
import os
from pathlib import Path
import tempfile

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXPECTED_TOOLS = {"get_status", "create_project", "list_projects", "import_source_text",
                  "read_source_text", "get_storyboard", "save_storyboard", "export_storyboard"}


async def check():
    entry = json.loads((ROOT / ".mcp.json").read_text(encoding="utf-8"))["mcpServers"]["shortdrama-basic"]
    if not Path(entry["command"]).is_absolute() or not Path(entry["command"]).is_file():
        raise ValueError("Run the installation/configuration script before this smoke test")
    with tempfile.TemporaryDirectory(prefix="shortdrama basic mcp ") as temporary:
        env = {**os.environ, **entry.get("env", {}), "SHORTDRAMA_BASIC_DATA_DIR": temporary}
        params = StdioServerParameters(command=entry["command"], args=entry["args"], cwd=entry["cwd"], env=env)
        async with stdio_client(params) as streams:
            async with ClientSession(*streams) as session:
                info = await session.initialize()
                available = await session.list_tools()
                assert {tool.name for tool in available.tools} == EXPECTED_TOOLS

                async def call(name, args=None):
                    response = await session.call_tool(name, args or {})
                    assert not response.is_error, f"{name}: {response.content}"
                    assert response.structured_content is not None
                    return response.structured_content

                status = await call("get_status")
                assert status["edition"] == "basic" and status["generation_engine"] is False
                project = await call("create_project", {"name": "基础版安装验证"})
                project_id = project["id"]
                projects = await call("list_projects")
                assert projects["total"] == 1
                original = "小雨在车站等失主。\r\n失主说：谢谢你。\n"
                source = await call("import_source_text", {"project_id": project_id, "title": "合成短文", "text": original})
                read = await call("read_source_text", {"project_id": project_id, "source_id": source["source_id"]})
                assert read["text"] == original
                board = await call("get_storyboard", {"project_id": project_id})
                shots = json.loads((ROOT / "examples/storyboard.json").read_text(encoding="utf-8"))
                for shot in shots:
                    shot["source_id"] = source["source_id"]
                saved = await call("save_storyboard", {"project_id": project_id, "shots": shots, "expected_revision": board["revision"]})
                stale = await session.call_tool("save_storyboard", {"project_id": project_id, "shots": shots, "expected_revision": board["revision"]})
                assert stale.is_error
                current = await call("get_storyboard", {"project_id": project_id})
                assert current["revision"] == saved["revision"] and current["shots"] == shots
                exported = await call("export_storyboard", {"project_id": project_id})
                path = Path(exported["path"])
                assert path.is_relative_to(Path(temporary)) and path.is_file()
                workbook = load_workbook(path)
                try:
                    assert workbook.sheetnames == ["Storyboard", "Project"]
                    assert workbook["Storyboard"].max_row == 3
                    assert workbook["Storyboard"]["E3"].value == "谢谢你。"
                finally:
                    workbook.close()
                return {"status": "pass", "version": info.server_info.version, "tools": sorted(EXPECTED_TOOLS),
                        "tool_count": len(EXPECTED_TOOLS), "source_roundtrip": "exact", "shots_saved": 2,
                        "revision_conflict_rejected": True, "xlsx_roundtrip": "pass",
                        "data_scope": "temporary synthetic data", "external_services_used": False}


if __name__ == "__main__":
    print(json.dumps(asyncio.run(check()), ensure_ascii=False, indent=2))
